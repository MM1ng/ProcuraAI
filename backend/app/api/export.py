from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel
from app.services.history_service import get_history_record
from app.services.confirmed_transaction import require_snapshot, SnapshotInvalid


router = APIRouter(prefix="/api/export", tags=["export"])


class ExcelExportRequest(BaseModel):
    model_config = {"extra": "forbid"}
    plan: dict[str, Any] | None = None
    history_id: str | None = None


HEADERS = [
    "Product",
    "Category",
    "Brand",
    "Supplier",
    "Quantity",
    "Unit Price",
    "Subtotal",
    "Rating",
    "Stock",
    "Delivery Days",
    "Why Selected",
    "Total Cost",
]


def build_procurement_workbook(plan: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Procurement Plan"
    sheet.append(HEADERS)

    total_cost = float(plan.get("total_amount", 0) or 0)
    items = plan.get("items") or plan.get("selected_items") or []
    for item in items:
        sheet.append(
            [
                item.get("name", ""),
                item.get("category", ""),
                item.get("brand", ""),
                item.get("supplier", ""),
                int(item.get("quantity", 0) or 0),
                float(item.get("unit_price", 0) or 0),
                float(item.get("subtotal", 0) or 0),
                float(item.get("rating", 0) or 0),
                int(item.get("stock", 0) or 0),
                int(item.get("delivery_days", 0) or 0),
                item.get("reason", ""),
                total_cost,
            ]
        )

    total_row = len(items) + 2
    sheet.cell(row=total_row, column=11, value="Total Cost")
    sheet.cell(row=total_row, column=12, value=total_cost)
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 48)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.post("/excel")
def export_excel(request: ExcelExportRequest) -> StreamingResponse:
    if request.history_id and request.plan is None:
        record = get_history_record(request.history_id)
        if not record:
            raise HTTPException(status_code=404, detail="History record not found")
        plan = record["selected_plan"]
    elif request.plan is not None and request.history_id is None:
        p = request.plan
        try:
            plan = require_snapshot(p.get("session_id"), p.get("plan_id"), p.get("plan_hash"), p.get("plan_version"))
        except SnapshotInvalid as exc:
            raise HTTPException(status_code=409, detail=exc.detail) from exc
    else:
        raise HTTPException(status_code=422, detail="Specify exactly one owned history or server plan.")
    output = build_procurement_workbook(plan)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="procurement_plan.xlsx"'},
    )
