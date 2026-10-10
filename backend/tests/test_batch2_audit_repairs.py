"""Audit-confirmed Batch 2 repairs, with isolated persistence and real SDK HMAC."""
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import hashlib
import hmac
import json
import sqlite3
import time
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
import stripe
from openpyxl import load_workbook
from test_caller_ownership_batch2 import browser, initialize, make_chat_order, ORIGIN
from test_transaction_snapshot_batch1 import ctx, plan as fixture_plan
import app.agent.procurement_agent as procurement_agent
import app.api.export as export_api
from app.services import caller_identity as identities, history_service, order_service, stripe_payment_service
import app.api.payments as payment_api

@pytest.fixture(autouse=True)
def isolated_history(monkeypatch, tmp_path):
    path = tmp_path / "history.json"
    path.write_text("[]", encoding="utf-8")
    monkeypatch.setattr(history_service, "HISTORY_FILE", path)

@contextmanager
def actor_context(client):
    actor, _ = identities.resolve_credential(client.cookies.get(identities.COOKIE_NAME))
    token = identities.CURRENT_CALLER.set(actor)
    try:
        yield actor
    finally:
        identities.CURRENT_CALLER.reset(token)

def mock_settings(monkeypatch):
    settings = SimpleNamespace(app_env="test", allow_mock_payment=True, use_mock_payment=True,
                               stripe_secret_key="", stripe_success_url="", stripe_cancel_url="", frontend_base_url=ORIGIN)
    monkeypatch.setattr(payment_api, "get_settings", lambda: settings)
    monkeypatch.setattr(stripe_payment_service, "get_settings", lambda: settings)

def checkout(a):
    _, p, _, _, order = make_chat_order(a)
    r = a.post("/api/payments/stripe/checkout", json={"order_id": order["order_id"]})
    assert r.status_code == 200, r.text
    return p, order, r.json()

@pytest.mark.parametrize("payment_mode", ["real", "mock"])
@pytest.mark.parametrize("status", ["pending_payment", "paid"])
def test_payment_return_is_readonly_and_displays_saved_status(ctx, monkeypatch, payment_mode, status):
    a = browser(); initialize(a)
    if payment_mode == "mock": mock_settings(monkeypatch)
    _, order, session = checkout(a)
    if status == "paid":
        with actor_context(a):
            order_service.update_order(order["order_id"], {"status": "paid"})
    before = order_service.ORDERS_FILE.read_bytes()
    write = Mock(wraps=order_service._write_orders)
    monkeypatch.setattr(order_service, "_write_orders", write)
    r = a.get("/api/payments/return-status", params={"order_id": order["order_id"], "session_id": session["session_id"]})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == status
    assert order_service.ORDERS_FILE.read_bytes() == before
    write.assert_not_called()

def test_payment_return_missing_foreign_and_forged_refs_have_no_effect(ctx):
    a, b = browser(), browser(); initialize(a); initialize(b)
    _, order, session = checkout(a)
    before = order_service.ORDERS_FILE.read_bytes()
    assert browser().get("/api/payments/return-status", params={"order_id": order["order_id"]}).status_code == 401
    assert b.get("/api/payments/return-status", params={"order_id": order["order_id"]}).status_code == 403
    assert a.get("/api/payments/return-status").status_code == 400
    assert a.get("/api/payments/return-status", params={"order_id": order["order_id"], "session_id": "forged"}).status_code == 400
    assert order_service.ORDERS_FILE.read_bytes() == before

@pytest.mark.parametrize("stored_checkout", [None, "cs_real_fixture", "mock_existing"])
def test_mock_action_requires_owned_existing_mock_checkout_and_exact_reference(ctx, monkeypatch, stored_checkout):
    mock_settings(monkeypatch)
    a = browser(); initialize(a)
    _, _, _, _, order = make_chat_order(a)
    with actor_context(a):
        order_service.update_order(order["order_id"], {"stripe_session_id": stored_checkout})
    before = order_service.ORDERS_FILE.read_bytes()
    # A query-string claim or order ID alone cannot advance payment.
    r = a.post("/api/payments/mock-success?mock=true", json={"order_id": order["order_id"]})
    assert r.status_code in (400, 403), r.text
    r = a.post("/api/payments/mock-success", json={"order_id": order["order_id"], "session_id": "forged"})
    assert r.status_code in (400, 403), r.text
    assert order_service.ORDERS_FILE.read_bytes() == before

def test_explicit_mock_action_bound_to_checkout_remains_available(ctx, monkeypatch):
    mock_settings(monkeypatch)
    a = browser(); initialize(a)
    _, order, session = checkout(a)
    r = a.post("/api/payments/mock-success", json={"order_id": order["order_id"], "session_id": session["session_id"]})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "paid"

def signed_event(order, **changes):
    obj = {"id": "cs_signed_repair", "object": "checkout.session", "metadata": {"order_id": order["order_id"]},
           "client_reference_id": order["order_id"], "payment_status": "paid", "amount_total": 2000, "currency": "usd"}
    obj.update(changes)
    return {"id": "evt_repair", "object": "event", "type": "checkout.session.completed", "data": {"object": obj}}

def post_signed(event):
    payload = json.dumps(event).encode()
    stamp = str(int(time.time()))
    digest = hmac.new(b"whsec_repair_fixture", stamp.encode() + b"." + payload, hashlib.sha256).hexdigest()
    return browser().post("/api/payments/stripe/webhook", content=payload,
                          headers={"stripe-signature": f"t={stamp},v1={digest}"})

@pytest.fixture
def signed_order(ctx, monkeypatch):
    a = browser(); initialize(a)
    _, _, _, _, order = make_chat_order(a)
    with actor_context(a):
        order_service.update_order(order["order_id"], {"stripe_session_id": "cs_signed_repair"})
    monkeypatch.setattr(stripe_payment_service, "get_settings", lambda: SimpleNamespace(stripe_webhook_secret="whsec_repair_fixture"))
    monkeypatch.setattr(stripe_payment_service, "_stripe_module", lambda: stripe)
    return order

@pytest.mark.parametrize("case", ["unpaid", "missing_amount", "wrong_amount", "bool_amount", "missing_currency",
                                 "wrong_currency", "wrong_checkout", "wrong_client_reference", "missing_order_references"])
def test_signed_invalid_payment_never_writes_paid(signed_order, monkeypatch, case):
    event = signed_event(signed_order)
    obj = event["data"]["object"]
    if case == "unpaid": obj["payment_status"] = "unpaid"
    if case == "missing_amount": obj.pop("amount_total")
    if case == "wrong_amount": obj["amount_total"] = 1
    if case == "bool_amount": obj["amount_total"] = True
    if case == "missing_currency": obj.pop("currency")
    if case == "wrong_currency": obj["currency"] = "cny"
    if case == "wrong_checkout": obj["id"] = "cs_foreign"
    if case == "wrong_client_reference": obj["client_reference_id"] = "OTHER_ORDER"
    if case == "missing_order_references": obj["metadata"] = {}; obj.pop("client_reference_id")
    before = order_service.ORDERS_FILE.read_bytes()
    write = Mock(wraps=order_service._write_orders)
    monkeypatch.setattr(order_service, "_write_orders", write)
    r = post_signed(event)
    assert r.status_code in (200, 400), r.text
    assert order_service.get_order(signed_order["order_id"])["status"] == "pending_payment"
    assert order_service.ORDERS_FILE.read_bytes() == before
    write.assert_not_called()

def test_signed_paid_replay_and_late_events_never_rewrite_or_downgrade(signed_order, monkeypatch):
    event = signed_event(signed_order)
    assert post_signed(event).status_code == 200
    before = order_service.ORDERS_FILE.read_bytes()
    write = Mock(wraps=order_service._write_orders)
    monkeypatch.setattr(order_service, "_write_orders", write)
    assert post_signed(event).status_code == 200
    for event_type in ("checkout.session.expired", "payment_intent.payment_failed"):
        late = deepcopy(event); late["id"] = "late-" + event_type; late["type"] = event_type
        assert post_signed(late).status_code == 200
    assert order_service.get_order(signed_order["order_id"])["status"] == "paid"
    assert order_service.ORDERS_FILE.read_bytes() == before
    write.assert_not_called()

def history_payload(p, order):
    return {"original_request": "Recommend 2 keyboards", "parsed_intent": {"categories": ["Keyboard"]},
            "procurement_plan": p, "trace": {}, "messages": [], "order_draft": order}

def test_history_owner_blocks_copies_and_all_existing_access_paths(ctx):
    a, b = browser(), browser(); initialize(a); initialize(b)
    _, p, _, _, order = make_chat_order(a)
    payload = history_payload(p, order)
    created = a.post("/api/history", json=payload)
    assert created.status_code == 200, created.text
    hid = created.json()["id"]
    assert b.get("/api/history").json()["items"] == []
    assert browser().get("/api/history").status_code == 401
    assert browser().get(f"/api/history/{hid}").status_code == 401
    assert b.get(f"/api/history/{hid}").status_code == 403
    assert b.delete(f"/api/history/{hid}").status_code == 403
    assert b.post("/api/history", json=payload).status_code == 403
    assert b.post("/api/history", json={**payload, "user_id": created.json().get("owner_actor_id")}).status_code == 422
    assert b.post("/api/export/excel", json={"history_id": hid}).status_code == 403
    assert browser().post("/api/export/excel", json={"history_id": hid}).status_code == 401
    assert b.post("/api/export/excel", json={"plan": p}).status_code == 403
    assert a.get(f"/api/history/{hid}").status_code == 200
    assert a.post("/api/export/excel", json={"history_id": hid}).status_code == 200
    assert a.delete(f"/api/history/{hid}").status_code == 200

def test_ownerless_history_is_not_adopted_or_exported(ctx):
    history_service.HISTORY_FILE.write_text(json.dumps([{"id": "LEGACY", "procurement_plan": {"items": []}, "order_draft": {"order_id": "OLD"}}]))
    a = browser(); initialize(a)
    before = history_service.HISTORY_FILE.read_bytes()
    assert a.get("/api/history").json()["items"] == []
    assert a.get("/api/history/LEGACY").status_code == 403
    assert a.delete("/api/history/LEGACY").status_code == 403
    assert a.post("/api/export/excel", json={"history_id": "LEGACY"}).status_code == 403
    assert history_service.HISTORY_FILE.read_bytes() == before

@pytest.mark.parametrize("invalid_kind", ["expired", "revoked", "forged"])
def test_explicit_invalid_cookie_clear_does_not_issue_actor_then_new_actor_cannot_claim_old(ctx, invalid_kind):
    a = browser(); initialize(a)
    _, _, _, _, order = make_chat_order(a)
    with actor_context(a) as actor:
        digest = actor.credential_digest
    if invalid_kind == "forged":
        a.cookies.clear(); a.cookies.set(identities.COOKIE_NAME, "forged", domain="127.0.0.1", path="/")
    else:
        with sqlite3.connect(identities.registry_path()) as db:
            column = "expires_at" if invalid_kind == "expired" else "revoked"
            db.execute(f"UPDATE actors SET {column}=? WHERE credential_digest=?", (0 if invalid_kind == "expired" else 1, digest))
    with sqlite3.connect(identities.registry_path()) as db:
        before_count = db.execute("SELECT count(*) FROM actors").fetchone()[0]
    before_orders = order_service.ORDERS_FILE.read_bytes()
    cleared = a.post("/api/caller/clear-invalid", headers={"X-Anonymous-Reset": "clear-invalid-context"})
    assert cleared.status_code == 200, cleared.text
    assert not a.cookies.get(identities.COOKIE_NAME)
    with sqlite3.connect(identities.registry_path()) as db:
        assert db.execute("SELECT count(*) FROM actors").fetchone()[0] == before_count
    assert order_service.ORDERS_FILE.read_bytes() == before_orders
    assert a.get("/api/caller/context").status_code == 401
    initialize(a)  # Separate explicit action, not part of the clear operation.
    assert a.get(f'/api/orders/{order["order_id"]}').status_code == 403
    assert a.get("/api/orders").json()["items"] == []

def test_clear_action_cannot_replace_valid_cookie(ctx):
    a = browser(); original = initialize(a)
    cookie = dict(a.cookies)
    r = a.post("/api/caller/clear-invalid", headers={"X-Anonymous-Reset": "clear-invalid-context"})
    assert r.status_code == 409 and "set-cookie" not in r.headers
    assert dict(a.cookies) == cookie
    assert a.get("/api/caller/context").json() == original.json()

@pytest.mark.parametrize("boundary", ["origin", "host", "marker"])
def test_invalid_cookie_clear_still_requires_explicit_loopback_boundary(ctx, boundary):
    a = browser()
    a.cookies.set(identities.COOKIE_NAME, "forged", domain="127.0.0.1", path="/")
    headers = {"X-Anonymous-Reset": "clear-invalid-context"}
    if boundary == "origin": headers["Origin"] = "http://localhost:3000"
    if boundary == "host": headers["Host"] = "localhost:8000"
    if boundary == "marker": headers["X-Anonymous-Reset"] = "automatic-reset"
    r = a.post("/api/caller/clear-invalid", headers=headers)
    assert r.status_code == 403 and "set-cookie" not in r.headers
    assert a.cookies.get(identities.COOKIE_NAME) == "forged"

def assert_archive_a_workbook(content):
    """Business oracle: A is p1 x2 at 10, total 20; B must never appear."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Procurement Plan"]
        rows = list(workbook.active.iter_rows(values_only=True))
        assert len(rows) == 3, "Expected exactly one A item and one total row"
        headers, item_row, total_row = rows
        exported = dict(zip(headers, item_row))
        # The existing format exposes Product names, not a separate ID column.
        # Fixture names are unique and identical to IDs, so p1 identifies A.
        for field, expected in {"Product": "p1", "Quantity": 2, "Unit Price": 10,
                                "Subtotal": 20, "Total Cost": 20}.items():
            assert exported[field] == expected, f"{field}: expected archived A value {expected}"
        assert total_row[-2] == "Total Cost"
        assert total_row[-1] == 20, "Summary total must be archived A total"
    finally:
        workbook.close()

@pytest.mark.parametrize("fault", ["none", "current_plan_B", "quantity", "total"])
def test_history_uses_server_plan_order_and_remains_exportable_as_owned_archive(ctx, monkeypatch, fault):
    a = browser(); initialize(a)
    _, p, _, _, order = make_chat_order(a)
    payload = history_payload(deepcopy(p), deepcopy(order))
    payload["procurement_plan"]["total_amount"] = 0.01
    payload["order_draft"]["total_amount"] = 0.01
    created = a.post("/api/history", json=payload)
    assert created.status_code == 200, created.text
    stored = created.json()
    assert stored["selected_plan"]["total_amount"] == p["total_amount"]
    assert stored["order_draft"]["total_amount"] == order["total_amount"]
    assert [(i["product_id"], i["name"], i["quantity"], i["unit_price"], i["subtotal"])
            for i in stored["selected_plan"]["items"]] == [("p1", "p1", 2, 10, 20)]

    # Force genuinely different content through the real agent/session registration.
    new_candidate = fixture_plan("p2", quantity=3)
    monkeypatch.setattr(procurement_agent, "parse_purchase_request", lambda *_: {
        "categories": ["Keyboard"], "budget": 100, "revision_intent": "new_plan",
        "quantity_by_category": {"Keyboard": 3}, "used_mock_llm": True})
    monkeypatch.setattr(procurement_agent, "generate_procurement_plan", lambda *_, **__: deepcopy(new_candidate))
    monkeypatch.setattr(procurement_agent, "generate_plan_options", lambda *_: [
        {"id": "new-b", "strategy": "premium", "plan": deepcopy(new_candidate)}])
    changed = a.post("/api/chat", json={"session_id": p["session_id"], "message": "Recommend 3 premium keyboards"})
    assert changed.status_code == 200, changed.text
    current = changed.json()["recommended_plan"]
    assert current["plan_id"] != p["plan_id"]
    assert current["plan_hash"] != p["plan_hash"]
    assert current["plan_version"] > p["plan_version"]
    assert [(i["product_id"], i["name"], i["quantity"], i["unit_price"], i["subtotal"])
            for i in current["items"]] == [("p2", "p2", 3, 15, 45)]
    assert current["total_amount"] == 45 != p["total_amount"]
    assert a.get(f'/api/history/{stored["id"]}').json()["selected_plan"] == stored["selected_plan"]

    # Supplemental fault injection retains the real owner lookup and XLSX writer.
    # It proves the content oracle catches wrong selection/quantity/amount at 200.
    writer = export_api.build_procurement_workbook
    if fault != "none":
        def corrupted_writer(selected):
            damaged = deepcopy(current if fault == "current_plan_B" else selected)
            if fault == "quantity": damaged["items"][0]["quantity"] = 3
            if fault == "total": damaged["total_amount"] = 99
            return writer(damaged)
        monkeypatch.setattr(export_api, "build_procurement_workbook", corrupted_writer)
    exported = a.post("/api/export/excel", json={"history_id": stored["id"]})
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if fault == "none":
        assert_archive_a_workbook(exported.content)
    else:
        expected_failure = {"current_plan_B": "Product", "quantity": "Quantity", "total": "Total Cost"}[fault]
        with pytest.raises(AssertionError, match=expected_failure):
            assert_archive_a_workbook(exported.content)
        monkeypatch.setattr(export_api, "build_procurement_workbook", writer)
        restored = a.post("/api/export/excel", json={"history_id": stored["id"]})
        assert restored.status_code == 200, restored.text
        assert_archive_a_workbook(restored.content)

def test_history_cannot_embed_another_actors_order_and_service_has_owner_guards(ctx):
    a, b = browser(), browser(); initialize(a); initialize(b)
    _, own_plan, _, _, own_order = make_chat_order(a)
    _, _, _, _, foreign_order = make_chat_order(b)
    assert a.post("/api/history", json=history_payload(own_plan, foreign_order)).status_code == 403
    stored = a.post("/api/history", json=history_payload(own_plan, own_order)).json()
    with actor_context(b):
        assert history_service.list_history_records() == []
        with pytest.raises(identities.CallerDenied):
            history_service.get_history_record(stored["id"])
        with pytest.raises(identities.CallerDenied):
            history_service.delete_history_record(stored["id"])

@pytest.mark.parametrize("bad_part", ["data", "object", "metadata", "event_id"])
def test_signed_malformed_evidence_rejects_without_writes(signed_order, bad_part):
    event = signed_event(signed_order)
    if bad_part == "data": event["data"] = ["invalid"]
    if bad_part == "object": event["data"]["object"] = ["invalid"]
    if bad_part == "metadata": event["data"]["object"]["metadata"] = ["invalid"]
    if bad_part == "event_id": event["id"] = None
    before = order_service.ORDERS_FILE.read_bytes()
    assert post_signed(event).status_code == 400
    assert order_service.ORDERS_FILE.read_bytes() == before

def test_signed_payment_without_a_bound_checkout_is_rejected(signed_order):
    records = json.loads(order_service.ORDERS_FILE.read_text())
    for order in records:
        if order["order_id"] == signed_order["order_id"]:
            order["stripe_session_id"] = None
    order_service.ORDERS_FILE.write_text(json.dumps(records), encoding="utf-8")
    event = signed_event(signed_order, id=None)
    before = order_service.ORDERS_FILE.read_bytes()
    assert post_signed(event).status_code == 400
    assert order_service.ORDERS_FILE.read_bytes() == before

@pytest.mark.parametrize("case", ["missing_amount", "missing_currency", "missing_references"])
def test_explicit_stripe_confirmation_requires_complete_payment_evidence(ctx, monkeypatch, case):
    a = browser(); initialize(a)
    _, order, session = checkout(a)
    evidence = {"id": session["session_id"], "payment_status": "paid", "amount_total": 2000, "currency": "usd",
                "metadata": {"order_id": order["order_id"]}, "client_reference_id": order["order_id"]}
    if case == "missing_amount": evidence.pop("amount_total")
    if case == "missing_currency": evidence.pop("currency")
    if case == "missing_references": evidence.pop("metadata"); evidence.pop("client_reference_id")
    monkeypatch.setattr(stripe_payment_service, "_stripe_module", lambda: SimpleNamespace(
        checkout=SimpleNamespace(Session=SimpleNamespace(retrieve=lambda *_, **__: evidence))))
    before = order_service.ORDERS_FILE.read_bytes()
    assert a.post("/api/payments/confirm", json={"order_id": order["order_id"], "session_id": session["session_id"]}).status_code == 400
    assert order_service.ORDERS_FILE.read_bytes() == before
